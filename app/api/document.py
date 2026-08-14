from fastapi import APIRouter, UploadFile, File, Form, status, HTTPException, Depends, Query
from datetime import datetime
import io
import pdfplumber
import docx
# DİKKAT: `langchain_text_splitters` bilerek modül düzeyinde import EDİLMİYOR
# (aşağıda upload_document içinde). Paketin __init__.py'si
from app.utils.logger import logger
from app.models.user import User
from app.services.chroma_service import get_collection
from app.services.auth_service import require_admin_role
from app.utils.dosya_dogrula import GENEL_RET, dosyayi_dogrula, max_upload_bayt

router = APIRouter(
    prefix="/document",
    tags=["Document Management"]
)

def format_table_to_markdown(table_data: list[list[str]]) -> str:
    if not table_data or not table_data[0]:
        return ""
    
    markdown_table = "\n"
    headers = [str(item).replace("\n", " ").strip() if item else "" for item in table_data[0]]
    markdown_table += "| " + " | ".join(headers) + " |\n"
    markdown_table += "|-" + "-|-".join(["-" * len(h) for h in headers]) + "-|\n"
    
    for row in table_data[1:]:
        row_values = [str(item).replace("\n", " ").strip() if item else "" for item in row]
        while len(row_values) < len(headers):
            row_values.append("")
        markdown_table += "| " + " | ".join(row_values[:len(headers)]) + " |\n"
    
    return markdown_table + "\n"

def process_pdf_content(file_bytes: bytes) -> str:
    extracted_text = ""
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                extracted_text += text + "\n\n"
            
            tables = page.extract_tables()
            for table in tables:
                extracted_text += format_table_to_markdown(table)
    return extracted_text

def process_docx_content(file_bytes: bytes) -> str:
    extracted_text = ""
    doc = docx.Document(io.BytesIO(file_bytes))
    
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            extracted_text += paragraph.text + "\n\n"
    
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            table_data.append(row_data)
        extracted_text += format_table_to_markdown(table_data)
        
    return extracted_text

@router.post("/upload", status_code=status.HTTP_201_CREATED)
async def upload_document(
    category: str = Form(...),
    file: UploadFile = File(...),
    current_user: User = Depends(require_admin_role)
):
    try:
        # Starlette UploadFile.size gövde çalışmadan önce dolu olabilir; 413'ü
        # 2 GB RAM tahsis etmeden önce vermek için read()'ten ÖNCE bakıyoruz.
        if file.size is not None and file.size > max_upload_bayt():
            logger.warning(
                f"Dosya reddedildi (boyut asimi, read oncesi): "
                f"dosya_adi={file.filename!r} boyut={file.size}"
            )
            raise HTTPException(
                status_code=status.HTTP_413_CONTENT_TOO_LARGE,
                detail="Dosya çok büyük",
            )

        file_content = await file.read()
        extracted_text = ""
        # Uzantı, boyut ve gerçek içerik imzası burada doğrulanıyor (Gün 21).
        file_extension = dosyayi_dogrula(file.filename, file_content)

        if file_extension == "pdf":
            try:
                extracted_text = process_pdf_content(file_content)
            except Exception as e:
                logger.error(f"PDF processing error: {str(e)}")
                # Ayrıntılı aşama mesajı K5'i deler; ayrım yalnızca log'da.
                raise HTTPException(status_code=400, detail=GENEL_RET)
        elif file_extension == "docx":
            try:
                extracted_text = process_docx_content(file_content)
            except Exception as e:
                logger.error(f"DOCX processing error: {str(e)}")
                raise HTTPException(status_code=400, detail=GENEL_RET)
        elif file_extension == "txt":
            try:
                extracted_text = file_content.decode("utf-8")
            except UnicodeDecodeError:
                raise HTTPException(status_code=400, detail=GENEL_RET)
        else:
            raise HTTPException(status_code=400, detail=GENEL_RET)

        if not extracted_text.strip():
            raise HTTPException(status_code=400, detail="Empty content")

        # Bölücü BURADA import ediliyor, modül düzeyinde değil: import zinciri
        # torch'a kadar iniyor ve uygulama açılışını ~25 saniye bekletiyordu.
        from langchain_text_splitters import RecursiveCharacterTextSplitter

        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            length_function=len,
            separators=["\n\n", "\n", " ", ""]
        )
        
        text_chunks = text_splitter.split_text(extracted_text)
        
        if not text_chunks:
            raise HTTPException(status_code=400, detail="Chunking failed")
            
        metadata_list = []
        id_list = []
        
        # .lower() YOK: Yanik.txt ile yanik.txt aynı chunk id'ye düşüp birbirinin
        # üzerine yazmasın (kaynak metadata'sı orijinal adı korur).
        safe_filename = file.filename.replace(" ", "_")
        current_date = datetime.now().strftime("%Y-%m-%d")
        
        for index, chunk in enumerate(text_chunks):
            metadata_list.append({
                "source": file.filename,
                "upload_date": current_date,
                "category": category,
                "chunk_index": index
            })
            id_list.append(f"{safe_filename}_chunk_{index}")
        
        koleksiyon = get_collection()

        # Eski chunk'ların id'leri ÖNCE okunuyor ama silme SONRAYA bırakılıyor:
        # önce silseydik, upsert patladığında önceki iyi sürüm de kaybolurdu.
        eski_kayitlar = koleksiyon.get(where={"source": file.filename})
        eski_idler = set(eski_kayitlar.get("ids") or [])

        koleksiyon.upsert(
            documents=text_chunks,
            metadatas=metadata_list,
            ids=id_list
        )

        # Yeni sürümde karşılığı olmayan eski chunk'lar siliniyor (hayalet chunk).
        # upsert yalnızca kendisine verilen id'lere dokunduğu için, daha kısa bir
        artakalan = sorted(eski_idler - set(id_list))
        if artakalan:
            koleksiyon.delete(ids=artakalan)

        logger.info(f"Uploaded and chunked: {file.filename} by {current_user.username}")
        
        return {
            "status": "success",
            "message": "Success",
            "total_chunks": len(text_chunks),
            "sample_metadata": metadata_list[0] if metadata_list else None
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Upload error: {str(e)}")
        raise HTTPException(status_code=500, detail="Upload error")


@router.get("/liste")
async def dokumanlari_listele(
    current_user: User = Depends(require_admin_role)
):
    """Yüklenen dokümanları kaynak dosya adına göre gruplayıp döndürür."""
    kayitlar = get_collection().get(include=["metadatas"])
    ustveriler = kayitlar.get("metadatas") or []

    # Dosya adı -> o dosyaya ait chunk'ların üstverileri
    gruplar: dict[str, list[dict]] = {}
    for ustveri in ustveriler:
        kaynak = (ustveri or {}).get("source")
        if kaynak is None:
            continue  # kaynağı olmayan kayıt listelenemez
        gruplar.setdefault(kaynak, []).append(ustveri)

    liste = []
    for kaynak, parcalar in gruplar.items():
        # Kategori ve tarih, chunk_index'i en küçük olan parçadan okunuyor:
        # üstveri tutarsız olsa bile çıktı rastgele değişmesin (K8).
        ilk = min(parcalar, key=lambda u: u.get("chunk_index", 0))
        liste.append({
            "kaynak": kaynak,
            "chunk_sayisi": len(parcalar),
            "kategori": ilk.get("category"),
            "yukleme_tarihi": ilk.get("upload_date"),
        })

    # Belirlenimci sıra (K4): ChromaDB get() dönüş sırasını garanti etmiyor.
    liste.sort(key=lambda kayit: kayit["kaynak"])
    return liste


@router.delete("")
async def dokumani_sil(
    kaynak: str = Query(..., min_length=1, description="Silinecek dosyanın adı"),
    current_user: User = Depends(require_admin_role)
):
    """Bir dosyaya ait bütün chunk'ları bilgi tabanından siler."""
    koleksiyon = get_collection()
    mevcut = koleksiyon.get(where={"source": kaynak})
    silinecek = mevcut.get("ids") or []

    if not silinecek:
        raise HTTPException(status_code=404, detail="Doküman bulunamadı")

    koleksiyon.delete(where={"source": kaynak})
    logger.info(
        f"Dokuman silindi: {kaynak} ({len(silinecek)} chunk) by {current_user.username}"
    )
    return {"silinen_chunk": len(silinecek)}
